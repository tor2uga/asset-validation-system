from fastapi import FastAPI, Depends, HTTPException
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy.exc import IntegrityError
from fastapi import Request, Form
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse
from app.core_templates import templates
from fastapi.staticfiles import StaticFiles
from app.services.auth_service import get_user_with_role_check
from app.database import engine, Base, SessionLocal
from app.services.attestation_service import process_attestation, check_attestation_readiness
from fastapi import UploadFile, File
from app.services.validation_service import submit_team_member_validation
from openpyxl import load_workbook
from pathlib import Path
from app.database import get_db

from openpyxl import Workbook
from app.pdf_utils import generate_attestation_pdf
from app import models, schemas
from app.automation import run_asset_workflow
from app.services.discovery_service import (
    submit_new_asset_discovery,
    approve_discovery_by_team_lead,
    reject_discovery_by_team_lead,
    approve_discovery_by_infosec,
    reject_discovery_by_infosec,
)
from app.services.final_asset_service import (
    get_final_assets_for_department_user,
    get_final_assets_for_admin,
    get_department_map,
    get_department_by_id,
)

app = FastAPI()
app.mount("/static", StaticFiles(directory="app/templates"), name="static")
app.add_middleware(SessionMiddleware, secret_key="super-secret-key-change-this")




ALLOWED_CONDITIONS = {"Good", "Decommisioned", "Faulty"}
ALLOWED_PRIVILEGES = {"Read-Only", "Local User", "Domain User", "Custodian", "Local Administrator"}
ALLOWED_CRITICALITIES = {"High", "Medium", "Low"}

@app.get("/login/{user_id}")
def login(user_id: int):
    return RedirectResponse(url=f"/ui/home/{user_id}", status_code=303)

# Create database tables from app/models.py
Base.metadata.create_all(bind=engine)


# -------------------------
# Database dependency
# -------------------------


#  CYcle dependency#

def get_active_cycle(db: Session):
    return db.query(models.ValidationCycle).filter(
        models.ValidationCycle.status == "active"
    ).order_by(models.ValidationCycle.start_date.desc()).first()
# -------------------------
# Root / health
# -------------------------
@app.get("/")
def read_root():
    return {"message": "Asset Automation API running 🚀"}


@app.post("/run-automation")
def run_automation():
    try:
        result = run_asset_workflow()
        return {"status": "success", "message": result}
    except Exception as e:
        return {"status": "error", "message": str(e)}


# -------------------------
# Departments
# -------------------------
@app.post("/departments/", response_model=schemas.DepartmentOut)
def create_department(dept: schemas.DepartmentCreate, db: Session = Depends(get_db)):
    new_dept = models.Department(name=dept.name)
    db.add(new_dept)

    try:
        db.commit()
        db.refresh(new_dept)
        return new_dept
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="Department already exists")


@app.get("/departments/", response_model=list[schemas.DepartmentOut])
def list_departments(db: Session = Depends(get_db)):
    return db.query(models.Department).order_by(models.Department.name).all()


# -------------------------
# AD Users (shadow users)
# -------------------------
@app.post("/ad-users/", response_model=schemas.ADUserOut)
def create_ad_user(user: schemas.ADUserCreate, db: Session = Depends(get_db)):
    new_user = models.ADUser(
        username=user.username,
        email=user.email,
        display_name=user.display_name
    )
    db.add(new_user)

    try:
        db.commit()
        db.refresh(new_user)
        return new_user
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="AD user already exists")


@app.get("/ad-users/", response_model=list[schemas.ADUserOut])
def list_ad_users(db: Session = Depends(get_db)):
    return db.query(models.ADUser).order_by(models.ADUser.username).all()


# -------------------------
# AD Group Mappings
# -------------------------
@app.post("/ad-group-mappings/", response_model=schemas.ADGroupMappingOut)
def create_ad_group_mapping(mapping: schemas.ADGroupMappingCreate, db: Session = Depends(get_db)):
    # Optional department existence check
    if mapping.department_id is not None:
        dept = db.query(models.Department).filter(models.Department.id == mapping.department_id).first()
        if not dept:
            raise HTTPException(status_code=404, detail="Department not found")

    new_mapping = models.ADGroupMapping(
        ad_group_name=mapping.ad_group_name,
        department_id=mapping.department_id,
        role=mapping.role
    )
    db.add(new_mapping)

    try:
        db.commit()
        db.refresh(new_mapping)
        return new_mapping
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="AD group mapping already exists")


@app.get("/ad-group-mappings/", response_model=list[schemas.ADGroupMappingOut])
def list_ad_group_mappings(db: Session = Depends(get_db)):
    return db.query(models.ADGroupMapping).order_by(models.ADGroupMapping.ad_group_name).all()


# -------------------------
# Validation Cycles
# -------------------------
@app.post("/validation-cycles/", response_model=schemas.ValidationCycleOut)
def create_validation_cycle(cycle: schemas.ValidationCycleCreate, db: Session = Depends(get_db)):
    new_cycle = models.ValidationCycle(
        name=cycle.name,
        start_date=cycle.start_date,
        end_date=cycle.end_date,
        status=cycle.status
    )
    db.add(new_cycle)

    try:
        db.commit()
        db.refresh(new_cycle)
        return new_cycle
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="Validation cycle already exists")


@app.get("/validation-cycles/", response_model=list[schemas.ValidationCycleOut])
def list_validation_cycles(db: Session = Depends(get_db)):
    return db.query(models.ValidationCycle).order_by(models.ValidationCycle.start_date.desc()).all()


# -------------------------
# Assets
# -------------------------
@app.post("/assets/", response_model=schemas.AssetOut)
def create_asset(asset: schemas.AssetCreate, db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(
        models.Department.id == asset.department_id
    ).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    new_asset = models.Asset(
        serial_no=asset.serial_no,
        asset_type=asset.asset_type,
        host_name=asset.host_name,
        description=asset.description,
        ip_address=asset.ip_address,
        privilege=asset.privilege,
        criticality=asset.criticality,
        department_id=asset.department_id
    )

    db.add(new_asset)
    db.commit()
    db.refresh(new_asset)
    return new_asset


@app.get("/assets/", response_model=list[schemas.AssetOut])
def list_assets(db: Session = Depends(get_db)):
    return db.query(models.Asset).order_by(models.Asset.id.desc()).all()


@app.get("/assets/department/{department_id}", response_model=list[schemas.AssetOut])
def get_assets_by_department(department_id: int, db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(
        models.Department.id == department_id
    ).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    return db.query(models.Asset).filter(
        models.Asset.department_id == department_id
    ).all()

# -------------------------
# Asset Validations
# -------------------------
@app.post("/asset-validations/", response_model=schemas.AssetValidationOut)
def create_asset_validation(data: schemas.AssetValidationCreate, db: Session = Depends(get_db)):
    # Check user
    user = db.query(models.ADUser).filter(models.ADUser.id == data.user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="AD user not found")

    # Check department
    dept = db.query(models.Department).filter(models.Department.id == data.department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    # Check cycle if provided
    if data.cycle_id is not None:
        cycle = db.query(models.ValidationCycle).filter(models.ValidationCycle.id == data.cycle_id).first()
        if not cycle:
            raise HTTPException(status_code=404, detail="Validation cycle not found")

    # Check asset if provided
    if data.asset_id is not None:
        asset = db.query(models.Asset).filter(models.Asset.id == data.asset_id).first()
        if not asset:
            raise HTTPException(status_code=404, detail="Asset not found")

    # Check reassignment department if provided
    if data.new_department_id is not None:
        new_dept = db.query(models.Department).filter(models.Department.id == data.new_department_id).first()
        if not new_dept:
            raise HTTPException(status_code=404, detail="New department not found")

    # Prevent duplicate validation for same asset/user/cycle when asset exists
    if data.asset_id is not None:
        existing = db.query(models.AssetValidation).filter(
            models.AssetValidation.asset_id == data.asset_id,
            models.AssetValidation.user_id == data.user_id,
            models.AssetValidation.cycle_id == data.cycle_id
        ).first()

        if existing:
            raise HTTPException(status_code=400, detail="Validation already exists for this asset/user/cycle")

    # Set automatic stage flags
    approval_stage = "pending_lead_review"
    reassignment_status = data.reassignment_status
    new_asset_status = data.new_asset_status

    if data.status == "reassign":
        reassignment_status = "pending_lead"

    new_validation = models.AssetValidation(
        cycle_id=data.cycle_id,
        asset_id=data.asset_id,
        user_id=data.user_id,
        department_id=data.department_id,
        status=data.status,
        comment=data.comment,
        is_new_asset=data.is_new_asset,
        new_asset_status=new_asset_status,
        reassignment_status=reassignment_status,
        new_department_id=data.new_department_id,
        approval_stage=approval_stage
    )

    db.add(new_validation)
    db.commit()
    db.refresh(new_validation)

    # Audit log
    audit = models.AuditLog(
        user_id=data.user_id,
        action="create_validation",
        entity_type="asset_validation",
        entity_id=new_validation.id,
        details=f"Status={data.status}; Department={data.department_id}; Asset={data.asset_id}"
    )
    db.add(audit)
    db.commit()

    return new_validation


@app.get("/asset-validations/", response_model=list[schemas.AssetValidationOut])
def list_asset_validations(db: Session = Depends(get_db)):
    return db.query(models.AssetValidation).order_by(models.AssetValidation.created_at.desc()).all()


@app.get("/asset-validations/department/{department_id}", response_model=list[schemas.AssetValidationOut])
def list_asset_validations_by_department(department_id: int, db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(models.Department.id == department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    return db.query(models.AssetValidation).filter(
        models.AssetValidation.department_id == department_id
    ).order_by(models.AssetValidation.created_at.desc()).all()


from datetime import datetime


# -------------------------
# Team Lead Approval
# -------------------------
@app.post("/asset-validations/{validation_id}/lead-approve", response_model=schemas.AssetValidationOut)
def lead_approve_validation(
    validation_id: int,
    lead_user_id: int,
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        raise HTTPException(status_code=404, detail="Validation not found")

    lead_user = db.query(models.ADUser).filter(
        models.ADUser.id == lead_user_id
    ).first()

    if not lead_user:
        raise HTTPException(status_code=404, detail="Lead user not found")

    if validation.approval_stage not in ["pending_lead_review", "returned_to_team"]:
        raise HTTPException(status_code=400, detail="Validation is not awaiting Team Lead review")

    validation.lead_approved_by = lead_user_id
    validation.lead_approved_at = datetime.utcnow()
    validation.approval_stage = "pending_attestation"

    if validation.status == "reassign":
        validation.reassignment_status = "pending_attestation"

    if validation.is_new_asset:
        validation.new_asset_status = "pending_attestation"

    db.commit()
    db.refresh(validation)

    audit = models.AuditLog(
        user_id=lead_user_id,
        action="lead_approve_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=f"approval_stage={validation.approval_stage}"
    )
    db.add(audit)
    db.commit()

    return validation


@app.post("/asset-validations/{validation_id}/lead-reject", response_model=schemas.AssetValidationOut)
def lead_reject_validation(
    validation_id: int,
    lead_user_id: int,
    reason: str,
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        raise HTTPException(status_code=404, detail="Validation not found")

    lead_user = db.query(models.ADUser).filter(
        models.ADUser.id == lead_user_id
    ).first()

    if not lead_user:
        raise HTTPException(status_code=404, detail="Lead user not found")

    validation.approval_stage = "returned_to_team"
    validation.comment = f"{validation.comment or ''} | Lead rejection reason: {reason}".strip()

    if validation.status == "reassign":
        validation.reassignment_status = "returned_to_team"

    if validation.is_new_asset:
        validation.new_asset_status = "returned_to_team"

    db.commit()
    db.refresh(validation)

    audit = models.AuditLog(
        user_id=lead_user_id,
        action="lead_reject_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=reason
    )
    db.add(audit)
    db.commit()

    return validation




@app.post("/team-attestations/", response_model=schemas.TeamAttestationOut)
def create_team_attestation(
    data: schemas.TeamAttestationCreate,
    db: Session = Depends(get_db)
):
    try:
        attestation = process_attestation(
            user_id=data.lead_user_id,
            department_id=data.department_id,
            cycle_id=data.cycle_id,
            attestation_data=data.attestation_data,
            db=db
        )
        return attestation
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))

    
@app.get("/team-attestations/", response_model=list[schemas.TeamAttestationOut])
def list_team_attestations(db: Session = Depends(get_db)):
    return db.query(models.TeamAttestation).order_by(
        models.TeamAttestation.submitted_at.desc()
    ).all()


@app.get("/team-attestations/{attestation_id}/download")
def download_attestation_pdf(
    attestation_id: int,
    user_id: int,
    db: Session = Depends(get_db)
):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if user.role != "compliance":
        raise HTTPException(status_code=403, detail="Only IT Compliance can download attestation PDFs")

    attestation = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.id == attestation_id
    ).first()
    if not attestation:
        raise HTTPException(status_code=404, detail="Attestation not found")

    if not attestation.file_path:
        raise HTTPException(status_code=404, detail="Attestation PDF not found")

    pdf_file = Path(attestation.file_path)
    if not pdf_file.exists():
        raise HTTPException(status_code=404, detail="Attestation PDF file missing")

    return FileResponse(
        path=str(pdf_file),
        filename=pdf_file.name,
        media_type="application/pdf"
    )



# -------------------------
# Infosec Approval
# Normal validations + new assets only
# -------------------------
@app.post("/asset-validations/{validation_id}/infosec-approve", response_model=schemas.AssetValidationOut)
def infosec_approve_validation(
    validation_id: int,
    infosec_user_id: int,
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        raise HTTPException(status_code=404, detail="Validation not found")

    infosec_user = db.query(models.ADUser).filter(
        models.ADUser.id == infosec_user_id
    ).first()

    if not infosec_user:
        raise HTTPException(status_code=404, detail="Infosec user not found")

    if validation.approval_stage != "pending_infosec":
        raise HTTPException(status_code=400, detail="Validation is not awaiting Infosec approval")

    if validation.status == "reassign":
        raise HTTPException(status_code=400, detail="Reassignment must be approved by IT Compliance, not Infosec")

    validation.infosec_approved_by = infosec_user_id
    validation.infosec_approved_at = datetime.utcnow()
    validation.approved = True
    validation.approval_stage = "approved"

    if validation.is_new_asset:
        validation.new_asset_status = "approved"

    db.commit()
    db.refresh(validation)

    audit = models.AuditLog(
        user_id=infosec_user_id,
        action="infosec_approve_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=f"status={validation.status}"
    )
    db.add(audit)
    db.commit()

    return validation


@app.post("/new-asset-discoveries/{discovery_id}/infosec-approve", response_model=schemas.NewAssetDiscoveryOut)
def infosec_approve_new_asset(
    discovery_id: int,
    infosec_user_id: int,
    db: Session = Depends(get_db)
):
    discovery = db.query(models.NewAssetDiscovery).filter(
        models.NewAssetDiscovery.id == discovery_id
    ).first()

    if not discovery:
        raise HTTPException(status_code=404, detail="New asset discovery not found")

    infosec_user = db.query(models.ADUser).filter(
        models.ADUser.id == infosec_user_id
    ).first()

    if not infosec_user:
        raise HTTPException(status_code=404, detail="Infosec user not found")

    if discovery.status != "new_asset_pending_infosec":
        raise HTTPException(status_code=400, detail="New asset is not awaiting Infosec approval")

    discovery.infosec_approved_by = infosec_user_id
    discovery.infosec_approved_at = datetime.utcnow()
    discovery.status = "approved"

    # Create official asset after Infosec approval
    new_asset = models.Asset(
        serial_no=discovery.asset_serial_no,
        asset_type=discovery.asset_type,
        host_name=discovery.host_application_server_name,
        description=discovery.description,
        ip_address=discovery.application_url_ip,
        privilege=discovery.privilege,
        criticality=discovery.criticality,
        department_id=discovery.department_id
    )
    db.add(new_asset)
    db.commit()
    db.refresh(discovery)

    audit = models.AuditLog(
        user_id=infosec_user_id,
        action="infosec_approve_new_asset",
        entity_type="new_asset_discovery",
        entity_id=discovery.id,
        details=f"department_id={discovery.department_id}; serial={discovery.asset_serial_no}"
    )
    db.add(audit)
    db.commit()

    return discovery


# -------------------------
# IT Compliance Approval
# Reassignment only
# -------------------------
@app.post("/ui/compliance/approve-reassignment")
def ui_compliance_approve_reassignment(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)

    if validation.status != "reassign":
        return HTMLResponse("Only reassignment validations are allowed here", status_code=400)

    if validation.approval_stage != "pending_it_compliance":
        return HTMLResponse("Validation is not awaiting IT Compliance approval", status_code=400)

    if not validation.new_department_id:
        return HTMLResponse("No target team found for reassignment", status_code=400)

    validation.compliance_approved_by = user_id
    validation.compliance_approved_at = datetime.utcnow()
    validation.reassignment_status = "approved"
    validation.approval_stage = "pending_new_team_validation"

    queue_item = models.TeamQueueItem(
        asset_id=validation.asset_id,
        department_id=validation.new_department_id,
        source_validation_id=validation.id,
        queue_type="reassigned_asset",
        status="pending_new_team_validation"
    )
    db.add(queue_item)

    notification = models.Notification(
        department_id=validation.new_department_id,
        title="Reassigned Asset Added to Queue",
        message="A reassigned asset has been approved and added to your team's validation queue."
    )
    db.add(notification)

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_compliance_approve_reassignment",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=f"new_department_id={validation.new_department_id}; queue_created_for_team={validation.new_department_id}"
    )
    db.add(audit)

    db.commit()

    return RedirectResponse(url=f"/ui/compliance/{user_id}", status_code=303)


@app.post("/ui/compliance/reject-reassignment")
def ui_compliance_reject_reassignment(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(""),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)

    validation.reassignment_status = "rejected"
    validation.approval_stage = "returned_to_lead"
    validation.comment = f"{validation.comment or ''} | Compliance rejection reason: {reason or 'No reason provided'}".strip()

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_compliance_reject_reassignment",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=reason or "No reason provided"
    )
    db.add(audit)

    db.commit()

    return RedirectResponse(url=f"/ui/compliance/{user_id}", status_code=303)

# -------------------------
# Audit Logs
# -------------------------
@app.get("/audit-logs/", response_model=list[schemas.AuditLogOut])
def list_audit_logs(db: Session = Depends(get_db)):
    return db.query(models.AuditLog).order_by(
        models.AuditLog.timestamp.desc()
    ).all()


@app.get("/audit-logs/entity/{entity_type}/{entity_id}", response_model=list[schemas.AuditLogOut])
def list_audit_logs_for_entity(entity_type: str, entity_id: int, db: Session = Depends(get_db)):
    return db.query(models.AuditLog).filter(
        models.AuditLog.entity_type == entity_type,
        models.AuditLog.entity_id == entity_id
    ).order_by(models.AuditLog.timestamp.desc()).all()

# -------------------------
# New Asset Discovery
# -------------------------
@app.post("/new-asset-discoveries/", response_model=schemas.NewAssetDiscoveryOut)
def create_new_asset_discovery(data: schemas.NewAssetDiscoveryCreate, db: Session = Depends(get_db)):
    # Check submitter
    user = db.query(models.ADUser).filter(models.ADUser.id == data.submitted_by).first()
    if not user:
        raise HTTPException(status_code=404, detail="Submitting user not found")

    # Check department
    dept = db.query(models.Department).filter(models.Department.id == data.department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    # Check cycle if provided
    if data.cycle_id is not None:
        cycle = db.query(models.ValidationCycle).filter(models.ValidationCycle.id == data.cycle_id).first()
        if not cycle:
            raise HTTPException(status_code=404, detail="Validation cycle not found")

    new_discovery = models.NewAssetDiscovery(
        cycle_id=data.cycle_id,
        department_id=data.department_id,
        submitted_by=data.submitted_by,
        sn=data.sn,
        condition=data.condition,
        asset_type=data.asset_type,
        asset_serial_no=data.asset_serial_no,
        host_application_server_name=data.host_application_server_name,
        description=data.description,
        application_url_ip=data.application_url_ip,
        privilege=data.privilege,
        criticality=data.criticality,
        assigned_to_custodian_owner_business_line=data.assigned_to_custodian_owner_business_line,
        status="new_asset_pending_lead_review"
    )

    db.add(new_discovery)
    db.commit()
    db.refresh(new_discovery)

    audit = models.AuditLog(
        user_id=data.submitted_by,
        action="create_new_asset_discovery",
        entity_type="new_asset_discovery",
        entity_id=new_discovery.id,
        details=f"Department={data.department_id}; Serial={data.asset_serial_no}"
    )
    db.add(audit)
    db.commit()

    return new_discovery


@app.get("/new-asset-discoveries/", response_model=list[schemas.NewAssetDiscoveryOut])
def list_new_asset_discoveries(db: Session = Depends(get_db)):
    return db.query(models.NewAssetDiscovery).order_by(models.NewAssetDiscovery.created_at.desc()).all()


@app.get("/new-asset-discoveries/department/{department_id}", response_model=list[schemas.NewAssetDiscoveryOut])
def list_new_asset_discoveries_by_department(department_id: int, db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(models.Department.id == department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    return db.query(models.NewAssetDiscovery).filter(
        models.NewAssetDiscovery.department_id == department_id
    ).order_by(models.NewAssetDiscovery.created_at.desc()).all()


# -------------------------
# Team Queue Items
# -------------------------
@app.post("/team-queue-items/", response_model=schemas.TeamQueueItemOut)
def create_team_queue_item(data: schemas.TeamQueueItemCreate, db: Session = Depends(get_db)):
    # Check asset
    asset = db.query(models.Asset).filter(models.Asset.id == data.asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    # Check department
    dept = db.query(models.Department).filter(models.Department.id == data.department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    # Check source validation if provided
    if data.source_validation_id is not None:
        validation = db.query(models.AssetValidation).filter(
            models.AssetValidation.id == data.source_validation_id
        ).first()
        if not validation:
            raise HTTPException(status_code=404, detail="Source validation not found")

    new_item = models.TeamQueueItem(
        asset_id=data.asset_id,
        department_id=data.department_id,
        source_validation_id=data.source_validation_id,
        queue_type=data.queue_type,
        status=data.status
    )

    db.add(new_item)
    db.commit()
    db.refresh(new_item)

    audit = models.AuditLog(
        user_id=None,
        action="create_team_queue_item",
        entity_type="team_queue_item",
        entity_id=new_item.id,
        details=f"Asset={data.asset_id}; Department={data.department_id}; QueueType={data.queue_type}"
    )
    db.add(audit)
    db.commit()

    return new_item


@app.get("/team-queue-items/", response_model=list[schemas.TeamQueueItemOut])
def list_team_queue_items(db: Session = Depends(get_db)):
    return db.query(models.TeamQueueItem).order_by(models.TeamQueueItem.created_at.desc()).all()


@app.get("/team-queue-items/department/{department_id}", response_model=list[schemas.TeamQueueItemOut])
def list_team_queue_items_by_department(department_id: int, db: Session = Depends(get_db)):
    dept = db.query(models.Department).filter(models.Department.id == department_id).first()
    if not dept:
        raise HTTPException(status_code=404, detail="Department not found")

    return db.query(models.TeamQueueItem).filter(
        models.TeamQueueItem.department_id == department_id
    ).order_by(models.TeamQueueItem.created_at.desc()).all()

@app.patch("/asset-validations/{validation_id}", response_model=schemas.AssetValidationOut)
def update_asset_validation(
    validation_id: int,
    data: schemas.AssetValidationCreate,
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        raise HTTPException(status_code=404, detail="Validation not found")

    # Optional guard: don't allow updates after it has moved beyond lead review
    if validation.approval_stage not in ["pending_lead_review", "returned_to_team"]:
        raise HTTPException(
            status_code=400,
            detail="Validation can no longer be edited at this stage"
        )

    # Check referenced entities if changed
    if data.user_id:
        user = db.query(models.ADUser).filter(models.ADUser.id == data.user_id).first()
        if not user:
            raise HTTPException(status_code=404, detail="AD user not found")

    if data.department_id:
        dept = db.query(models.Department).filter(models.Department.id == data.department_id).first()
        if not dept:
            raise HTTPException(status_code=404, detail="Department not found")

    if data.cycle_id is not None:
        cycle = db.query(models.ValidationCycle).filter(models.ValidationCycle.id == data.cycle_id).first()
        if not cycle:
            raise HTTPException(status_code=404, detail="Validation cycle not found")

    if data.asset_id is not None:
        asset = db.query(models.Asset).filter(models.Asset.id == data.asset_id).first()
        if not asset:
            raise HTTPException(status_code=404, detail="Asset not found")

    if data.new_department_id is not None:
        new_dept = db.query(models.Department).filter(models.Department.id == data.new_department_id).first()
        if not new_dept:
            raise HTTPException(status_code=404, detail="New department not found")

    old_status = validation.status
    old_new_department_id = validation.new_department_id
    old_comment = validation.comment

    # Apply updates
    validation.cycle_id = data.cycle_id
    validation.asset_id = data.asset_id
    validation.user_id = data.user_id
    validation.department_id = data.department_id
    validation.status = data.status
    validation.comment = data.comment
    validation.is_new_asset = data.is_new_asset
    validation.new_department_id = data.new_department_id

    # Reset derived statuses based on latest change
    validation.approval_stage = "pending_lead_review"

    if data.status == "reassign":
        validation.reassignment_status = "pending_lead"
    else:
        validation.reassignment_status = None

    if data.is_new_asset:
        validation.new_asset_status = "pending_lead"
    else:
        validation.new_asset_status = None

    # Clear approval markers because the record changed
    validation.approved = False
    validation.lead_approved_by = None
    validation.lead_approved_at = None
    validation.infosec_approved_by = None
    validation.infosec_approved_at = None
    validation.compliance_approved_by = None
    validation.compliance_approved_at = None

    db.commit()
    db.refresh(validation)

    audit = models.AuditLog(
        user_id=data.user_id,
        action="update_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=(
            f"old_status={old_status}, new_status={validation.status}; "
            f"old_new_department_id={old_new_department_id}, new_new_department_id={validation.new_department_id}; "
            f"old_comment={old_comment}, new_comment={validation.comment}"
        )
    )
    db.add(audit)
    db.commit()

    return validation

# -------------------------
# Team Member UI
# -------------------------
@app.get("/ui/team-member/{user_id}", response_class=HTMLResponse)
def team_member_dashboard(user_id: int, request: Request, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["team_member"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)

    if not user.department_id:
        return HTMLResponse("No department assigned", status_code=400)

    active_cycle = get_active_cycle(db)

    notifications = db.query(models.Notification).filter(
        models.Notification.department_id == user.department_id,
        models.Notification.is_read == False
    ).all()

    assets = db.query(models.Asset).filter(
        models.Asset.department_id == user.department_id
    ).order_by(models.Asset.id.desc()).all()

    departments = db.query(models.Department).order_by(models.Department.name).all()

    validations = db.query(models.AssetValidation).filter(
        models.AssetValidation.user_id == user_id
    ).all()

    validated_asset_ids = {
        v.asset_id for v in validations
        if v.asset_id is not None
        and v.cycle_id == (active_cycle.id if active_cycle else None)
        and v.approval_stage not in ["returned_to_team"]
    }

    returned_validation_map = {
        v.asset_id: v for v in validations
        if v.asset_id is not None
        and v.approval_stage == "returned_to_team"
    }

    return templates.TemplateResponse(
    request,
    "team_member_dashboard.html",
    {
        "request": request,
        "user": user,
        "assets": assets,
        "departments": departments,
        "validated_asset_ids": validated_asset_ids,
        "returned_validation_map": returned_validation_map,
        "notifications": notifications,
        "active_cycle": active_cycle
    }
)

@app.post("/ui/team-member/validate")
def ui_submit_validation(
    user_id: int = Form(...),
    asset_id: int = Form(...),
    status: str = Form(...),
    comment: str = Form(""),
    new_department_id: str = Form(None),
    db: Session = Depends(get_db)
):
    try:
        submit_team_member_validation(
            user_id=user_id,
            asset_id=asset_id,
            status=status,
            comment=comment,
            new_department_id=new_department_id,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/team-member/{user_id}", status_code=303)


@app.post("/ui/team-member/new-asset")
def ui_create_new_asset(
    user_id: int = Form(...),
    sn: str = Form(""),
    condition: str = Form(""),
    asset_type: str = Form(""),
    asset_serial_no: str = Form(""),
    host_application_server_name: str = Form(""),
    description: str = Form(""),
    application_url_ip: str = Form(""),
    privilege: str = Form(""),
    criticality: str = Form(""),
    assigned_to_custodian_owner_business_line: str = Form(""),
    db: Session = Depends(get_db)
):
    try:
        submit_new_asset_discovery(
            user_id=user_id,
            sn=sn,
            condition=condition,
            asset_type=asset_type,
            asset_serial_no=asset_serial_no,
            host_application_server_name=host_application_server_name,
            description=description,
            application_url_ip=application_url_ip,
            privilege=privilege,
            criticality=criticality,
            assigned_to_custodian_owner_business_line=assigned_to_custodian_owner_business_line,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/team-member/{user_id}", status_code=303)
    


# -------------------------
# Team Lead UI
# -------------------------
@app.get("/ui/team-lead/{user_id}", response_class=HTMLResponse)
def team_lead_dashboard(user_id: int, request: Request, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["team_lead"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)

    if not user.department_id:
        return HTMLResponse("No department assigned", status_code=400)

    notifications = db.query(models.Notification).filter(
        models.Notification.department_id == user.department_id,
        models.Notification.is_read == False
    ).all()

    validations = db.query(models.AssetValidation).filter(
    models.AssetValidation.department_id == user.department_id,
    models.AssetValidation.approval_stage.in_([
        "pending_lead_review",
        "returned_to_lead"
    ])
).all()

    for v in validations:
        v.asset = db.query(models.Asset).filter(
            models.Asset.id == v.asset_id
        ).first()

    discoveries = db.query(models.NewAssetDiscovery).filter(
        models.NewAssetDiscovery.department_id == user.department_id,
        models.NewAssetDiscovery.status.in_([
            "new_asset_pending_lead_review",
            "returned_to_lead"
        ])
    ).all()

    departments = db.query(models.Department).filter(
        models.Department.id == user.department_id
    ).all()

    cycles = db.query(models.ValidationCycle).all()

    existing_attestations = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.department_id == user.department_id
    ).all()

    attested_cycle_ids = {
        a.cycle_id for a in existing_attestations if a.cycle_id is not None
    }

    # ✅ STEP 3 GOES HERE
    attestation_readiness = {}

    for c in cycles:
        attestation_readiness[c.id] = check_attestation_readiness(
            department_id=user.department_id,
            cycle_id=c.id,
            db=db
        )

    return templates.TemplateResponse(
    request,
    "team_lead_dashboard.html",
    {
        "request": request,
        "user": user,
        "validations": validations,
        "discoveries": discoveries,
        "departments": departments,
        "cycles": cycles,
        "notifications": notifications,
        "attested_cycle_ids": attested_cycle_ids,
        "attestation_readiness": attestation_readiness
    }
)
    
@app.post("/ui/team-lead/approve-validation")
def ui_team_lead_approve_validation(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)
    
    if validation.approval_stage != "pending_lead_review":
        return HTMLResponse("This validation has already been processed.", status_code=400)

    if validation.status == "reassign":
        validation.approval_stage = "pending_attestation"
        validation.reassignment_status = "pending_attestation"
    else:
        validation.approval_stage = "pending_attestation"

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_team_lead_approve_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details="Approved by Team Lead"
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)


@app.post("/ui/team-lead/reject-validation")
def ui_team_lead_reject_validation(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(""),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)
    if validation.approval_stage != "pending_lead_review":
        return HTMLResponse("This validation has already been processed.", status_code=400)

    validation.approval_stage = "returned_to_team"
    validation.comment = f"{validation.comment or ''} | Lead rejection reason: {reason or 'No reason provided'}".strip()

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_team_lead_reject_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=reason or "No reason provided"
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)

@app.post("/ui/team-lead/approve-discovery")
def ui_team_lead_approve_discovery(
    discovery_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    try:
        approve_discovery_by_team_lead(
            discovery_id=discovery_id,
            user_id=user_id,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)


@app.post("/ui/team-lead/reject-discovery")
def ui_team_lead_reject_discovery(
    discovery_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(""),
    db: Session = Depends(get_db)
):
    try:
        reject_discovery_by_team_lead(
            discovery_id=discovery_id,
            user_id=user_id,
            reason=reason,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)

@app.post("/ui/team-lead/attest")
def ui_attestation(
    user_id: int = Form(...),
    department_id: int = Form(...),
    cycle_id: int = Form(...),
    attestation_data: str = Form(...),
    db: Session = Depends(get_db)
):
    try:
        process_attestation(
            user_id=user_id,
            department_id=department_id,
            cycle_id=cycle_id,
            attestation_data=attestation_data,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)
    

# -------------------------
# IT Compliance UI
# -------------------------
@app.get("/ui/compliance/{user_id}", response_class=HTMLResponse)
def compliance_dashboard(user_id: int, request: Request, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["compliance"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)

    # ✅ FIX: always define notifications
    notifications = []

    validations = db.query(models.AssetValidation).filter(
        models.AssetValidation.approval_stage == "pending_it_compliance",
        models.AssetValidation.status == "reassign"
    ).order_by(models.AssetValidation.created_at.desc()).all()

    departments = {
        d.id: d.name for d in db.query(models.Department).all()
    }
    assets = {
        a.id: a for a in db.query(models.Asset).all()
    }
    users = {
        u.id: u for u in db.query(models.ADUser).all()
    }

    attestations = db.query(models.TeamAttestation).order_by(
        models.TeamAttestation.submitted_at.desc()
    ).all()

    enriched_validations = []
    for v in validations:
        asset = assets.get(v.asset_id)
        request_user = users.get(v.user_id)
        current_team = departments.get(v.department_id, "N/A")
        new_team = departments.get(v.new_department_id, "N/A") if v.new_department_id else "-"

        enriched_validations.append({
            "id": v.id,
            "asset_id": v.asset_id,
            "asset_serial": asset.serial_no if asset else "N/A",
            "asset_type": asset.asset_type if asset else "N/A",
            "host_name": asset.host_name if asset else "N/A",
            "requested_by": request_user.display_name if request_user and request_user.display_name else (
                request_user.username if request_user else "N/A"
            ),
            "current_team": current_team,
            "new_team": new_team,
            "comment": v.comment,
            "approval_stage": v.approval_stage,
            "created_at": v.created_at
        })

    return templates.TemplateResponse(
    request,
    "compliance_dashboard.html",
    {
        "request": request,
        "user": user,
        "validations": enriched_validations,
        "attestations": attestations,
        "departments": departments,
        "users": users,
        "notifications": notifications
    }
)
# -------------------------
# Infosec UI
# -------------------------
@app.get("/ui/infosec/{user_id}", response_class=HTMLResponse)
def infosec_dashboard(user_id: int, request: Request, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["infosec"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)
    
    if user.role != "infosec":
        return HTMLResponse("Access denied", status_code=403)

    notifications = []

    departments = {d.id: d.name for d in db.query(models.Department).all()}
    users = {u.id: u for u in db.query(models.ADUser).all()}
    cycles = {c.id: c for c in db.query(models.ValidationCycle).all()}

    attestations = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.infosec_status == "pending"
    ).order_by(models.TeamAttestation.submitted_at.desc()).all()

    return templates.TemplateResponse(
    request,
    "infosec_dashboard.html",
    {
        "request": request,
        "user": user,
        "attestations": attestations,
        "departments": departments,
        "users": users,
        "cycles": cycles,
        "notifications": notifications
    }
)
    
@app.post("/ui/infosec/approve-validation")
def ui_infosec_approve_validation(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)

    if validation.approval_stage != "pending_infosec":
        return HTMLResponse("Validation is not awaiting Infosec approval", status_code=400)

    if validation.status == "reassign":
        return HTMLResponse("Reassignment does not belong in Infosec queue", status_code=400)

    validation.infosec_approved_by = user_id
    validation.infosec_approved_at = datetime.utcnow()
    validation.approved = True
    validation.approval_stage = "approved"

    if validation.status == "decommissioned":
        asset = db.query(models.Asset).filter(models.Asset.id == validation.asset_id).first()
        if asset:
            asset.description = f"{asset.description or ''} | DECOMMISSIONED".strip()

    if validation.status == "inactive":
        asset = db.query(models.Asset).filter(models.Asset.id == validation.asset_id).first()
        if asset:
            asset.description = f"{asset.description or ''} | INACTIVE".strip()

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_infosec_approve_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=f"status={validation.status}"
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

@app.post("/ui/infosec/reject-validation")
def ui_infosec_reject_validation(
    validation_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(...),
    db: Session = Depends(get_db)
):
    validation = db.query(models.AssetValidation).filter(
        models.AssetValidation.id == validation_id
    ).first()

    if not validation:
        return HTMLResponse("Validation not found", status_code=404)

    validation.approval_stage = "returned_to_lead"
    validation.comment = f"{validation.comment or ''} | Infosec rejection reason: {reason}".strip()

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_infosec_reject_validation",
        entity_type="asset_validation",
        entity_id=validation.id,
        details=reason
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

@app.post("/ui/infosec/approve-new-asset")
def ui_infosec_approve_new_asset(
    discovery_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    try:
        approve_discovery_by_infosec(
            discovery_id=discovery_id,
            user_id=user_id,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

@app.post("/ui/infosec/reject-new-asset")
def ui_infosec_reject_new_asset(
    discovery_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(""),
    db: Session = Depends(get_db)
):
    try:
        reject_discovery_by_infosec(
            discovery_id=discovery_id,
            user_id=user_id,
            reason=reason,
            db=db
        )
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)


@app.post("/ui/infosec/approve-attestation")
def ui_infosec_approve_attestation(
    attestation_id: int = Form(...),
    user_id: int = Form(...),
    db: Session = Depends(get_db)
):
    attestation = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.id == attestation_id
    ).first()

    if not attestation:
        return HTMLResponse("Attestation not found", status_code=404)

    attestation.infosec_status = "approved"
    attestation.infosec_approved_by = user_id
    attestation.infosec_approved_at = datetime.utcnow()
    attestation.infosec_comment = "Approved by Infosec"

    validations = db.query(models.AssetValidation).filter(
        models.AssetValidation.department_id == attestation.department_id,
        models.AssetValidation.cycle_id == attestation.cycle_id,
        models.AssetValidation.approval_stage == "awaiting_infosec_attestation"
    ).all()

    for v in validations:
        v.approval_stage = "approved"

    discoveries = db.query(models.NewAssetDiscovery).filter(
        models.NewAssetDiscovery.department_id == attestation.department_id,
        models.NewAssetDiscovery.cycle_id == attestation.cycle_id,
        models.NewAssetDiscovery.status == "awaiting_infosec_attestation"
    ).all()

    for d in discoveries:
        new_asset = models.Asset(
            serial_no=d.sn or d.asset_serial_no or f"new-{d.id}",
            asset_serial_no=d.asset_serial_no,
            asset_type=d.asset_type,
            host_name=d.host_application_server_name,
            description=d.description,
            ip_address=d.application_url_ip,
            privilege=d.privilege,
            criticality=d.criticality,
            condition=d.condition,
            assigned_to_custodian_owner_business_line=d.assigned_to_custodian_owner_business_line,
            department_id=d.department_id
        )
        db.add(new_asset)

        d.status = "approved"
        d.infosec_approved_by = user_id
        d.infosec_approved_at = datetime.utcnow()

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_infosec_approve_attestation",
        entity_type="team_attestation",
        entity_id=attestation.id,
        details=f"department_id={attestation.department_id}; cycle_id={attestation.cycle_id}"
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

@app.post("/ui/infosec/reject-attestation")
def ui_infosec_reject_attestation(
    attestation_id: int = Form(...),
    user_id: int = Form(...),
    reason: str = Form(""),
    db: Session = Depends(get_db)
):
    attestation = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.id == attestation_id
    ).first()

    if not attestation:
        return HTMLResponse("Attestation not found", status_code=404)

    attestation.infosec_status = "rejected"
    attestation.infosec_comment = reason or "No reason provided"

    validations = db.query(models.AssetValidation).filter(
        models.AssetValidation.department_id == attestation.department_id,
        models.AssetValidation.cycle_id == attestation.cycle_id,
        models.AssetValidation.approval_stage == "awaiting_infosec_attestation"
    ).all()

    for v in validations:
        v.approval_stage = "returned_to_lead"
        v.comment = f"{v.comment or ''} | Infosec attestation rejection: {reason or 'No reason provided'}".strip()

    discoveries = db.query(models.NewAssetDiscovery).filter(
        models.NewAssetDiscovery.department_id == attestation.department_id,
        models.NewAssetDiscovery.cycle_id == attestation.cycle_id,
        models.NewAssetDiscovery.status == "awaiting_infosec_attestation"
    ).all()

    for d in discoveries:
        d.status = "returned_to_lead"

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_infosec_reject_attestation",
        entity_type="team_attestation",
        entity_id=attestation.id,
        details=reason or "No reason provided"
    )
    db.add(audit)
    db.commit()

    return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

# -------------------------
# New Team Queue UI
# -------------------------
@app.get("/ui/new-team-queue/{user_id}", response_class=HTMLResponse)
def new_team_queue_dashboard(user_id: int, request: Request, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        return HTMLResponse("User not found", status_code=404)

    if user.role != "team_member":
        return HTMLResponse("Access denied", status_code=403)

    if not user.department_id:
        return HTMLResponse("No department assigned", status_code=400)

    queue_items = db.query(models.TeamQueueItem).filter(
        models.TeamQueueItem.department_id == user.department_id,
        models.TeamQueueItem.status == "pending_new_team_validation"
    ).order_by(models.TeamQueueItem.created_at.desc()).all()

    assets = {a.id: a for a in db.query(models.Asset).all()}
    departments = {d.id: d.name for d in db.query(models.Department).all()}
    validations = {v.id: v for v in db.query(models.AssetValidation).all()}
    cycles = db.query(models.ValidationCycle).order_by(models.ValidationCycle.start_date.desc()).all()

    notifications = db.query(models.Notification).filter(
        models.Notification.department_id == user.department_id,
        models.Notification.is_read == False
    ).order_by(models.Notification.created_at.desc()).all()

    enriched_queue = []
    for q in queue_items:
        asset = assets.get(q.asset_id)
        source_validation = validations.get(q.source_validation_id)

        enriched_queue.append({
            "queue_id": q.id,
            "asset_id": q.asset_id,
            "asset_serial": asset.serial_no if asset else "N/A",
            "asset_type": asset.asset_type if asset else "N/A",
            "host_name": asset.host_name if asset else "N/A",
            "description": asset.description if asset else "N/A",
            "ip_address": asset.ip_address if asset else "N/A",
            "privilege": asset.privilege if asset else "N/A",
            "criticality": asset.criticality if asset else "N/A",
            "team_name": departments.get(q.department_id, "N/A"),
            "source_validation_id": q.source_validation_id,
            "source_comment": source_validation.comment if source_validation else "",
            "department_id": q.department_id
        })

    return templates.TemplateResponse(
    request,
    "new_team_queue_dashboard.html",
    {
        "request": request,
        "user": user,
        "queue_items": enriched_queue,
        "cycles": cycles,
        "notifications": notifications
    }
)

@app.post("/ui/new-team-queue/validate")
def ui_validate_reassigned_queue_item(
    queue_id: int = Form(...),
    user_id: int = Form(...),
    cycle_id: int = Form(...),
    asset_id: int = Form(...),
    department_id: int = Form(...),
    status: str = Form(...),
    comment: str = Form(""),
    db: Session = Depends(get_db)
):
    queue_item = db.query(models.TeamQueueItem).filter(
        models.TeamQueueItem.id == queue_id
    ).first()

    if not queue_item:
        return HTMLResponse("Queue item not found", status_code=404)

    existing = db.query(models.AssetValidation).filter(
        models.AssetValidation.user_id == user_id,
        models.AssetValidation.asset_id == asset_id,
        models.AssetValidation.cycle_id == cycle_id,
        models.AssetValidation.department_id == department_id
    ).first()

    if existing:
        return RedirectResponse(url=f"/ui/new-team-queue/{user_id}", status_code=303)

    validation = models.AssetValidation(
        cycle_id=cycle_id,
        asset_id=asset_id,
        user_id=user_id,
        department_id=department_id,
        status=status,
        comment=comment,
        approval_stage="pending_lead_review"
    )
    db.add(validation)

    queue_item.status = "submitted_by_new_team"

    audit = models.AuditLog(
        user_id=user_id,
        action="ui_validate_reassigned_queue_item",
        entity_type="team_queue_item",
        entity_id=queue_item.id,
        details=f"asset_id={asset_id}; department_id={department_id}; status={status}"
    )
    db.add(audit)

    db.commit()

    return RedirectResponse(url=f"/ui/new-team-queue/{user_id}", status_code=303)

@app.get("/notifications/department/{department_id}", response_model=list[schemas.NotificationOut])
def get_department_notifications(department_id: int, db: Session = Depends(get_db)):
    return db.query(models.Notification).filter(
        models.Notification.department_id == department_id
    ).order_by(models.Notification.created_at.desc()).all()
    
@app.post("/notifications/mark-read/{notification_id}")
def mark_notification_read(notification_id: int, db: Session = Depends(get_db)):
    notification = db.query(models.Notification).filter(
        models.Notification.id == notification_id
    ).first()

    if notification:
        notification.is_read = True
        db.commit()

    return {"status": "ok"}


@app.get("/ui/home/{user_id}")
def route_user_dashboard(user_id: int, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()

    if not user:
        return HTMLResponse("User not found", status_code=404)

    if user.role == "team_member":
        return RedirectResponse(url=f"/ui/team-member/{user_id}", status_code=303)

    elif user.role == "team_lead":
        return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)

    elif user.role == "compliance":
        return RedirectResponse(url=f"/ui/compliance/{user_id}", status_code=303)

    elif user.role == "infosec":
        return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

    else:
        return HTMLResponse("No role assigned", status_code=400)
    
@app.get("/login/{user_id}")
def login(user_id: int):
    return RedirectResponse(url=f"/ui/home/{user_id}", status_code=303)

@app.get("/ui/home/{user_id}")
def route_user_dashboard(user_id: int, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()

    if not user:
        return HTMLResponse("User not found", status_code=404)

    if user.role == "team_member":
        return RedirectResponse(url=f"/ui/team-member/{user_id}", status_code=303)
    elif user.role == "team_lead":
        return RedirectResponse(url=f"/ui/team-lead/{user_id}", status_code=303)
    elif user.role == "compliance":
        return RedirectResponse(url=f"/ui/compliance/{user_id}", status_code=303)
    elif user.role == "infosec":
        return RedirectResponse(url=f"/ui/infosec/{user_id}", status_code=303)

    return HTMLResponse("No role assigned", status_code=400)


@app.post("/ad-users/", response_model=schemas.ADUserOut)
def create_ad_user(user: schemas.ADUserCreate, db: Session = Depends(get_db)):
    new_user = models.ADUser(
        username=user.username,
        email=user.email,
        display_name=user.display_name,
        role=user.role,
        department_id=user.department_id
    )
    db.add(new_user)

    try:
        db.commit()
        db.refresh(new_user)
        return new_user
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail="AD user already exists")
    
    
@app.patch("/ad-users/{user_id}/role", response_model=schemas.ADUserOut)
def update_ad_user_role(
    user_id: int,
    role: str,
    department_id: int | None = None,
    db: Session = Depends(get_db)
):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    allowed_roles = {"team_member", "team_lead", "compliance", "infosec"}
    if role not in allowed_roles:
        raise HTTPException(status_code=400, detail="Invalid role")

    if department_id is not None:
        dept = db.query(models.Department).filter(models.Department.id == department_id).first()
        if not dept:
            raise HTTPException(status_code=404, detail="Department not found")
        user.department_id = department_id

    user.role = role
    db.commit()
    db.refresh(user)
    return user

@app.get("/team-attestations/{attestation_id}/download")
def download_attestation_pdf(
    attestation_id: int,
    user_id: int,
    db: Session = Depends(get_db)
):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")
    
    print("DOWNLOAD DEBUG:", user_id, user.username if user else None, user.role if user else None)
    
    if user.role != "compliance":
        raise HTTPException(status_code=403, detail="Only IT Compliance can download attestation PDFs")
    
    attestation = db.query(models.TeamAttestation).filter(
        models.TeamAttestation.id == attestation_id
    ).first()
    if not attestation:
        raise HTTPException(status_code=404, detail="Attestation not found")

    if not attestation.file_path:
        raise HTTPException(status_code=404, detail="Attestation PDF not found")

    pdf_file = Path(attestation.file_path)
    if not pdf_file.exists():
        raise HTTPException(status_code=404, detail="Attestation PDF file missing")

    return FileResponse(
        path=str(pdf_file),
        filename=pdf_file.name,
        media_type="application/pdf"
    )
    

@app.get("/ui/{role}/{user_id}/final-assets", response_class=HTMLResponse)
def final_assets_dashboard(role: str, user_id: int, request: Request, db: Session = Depends(get_db)):
    normalized_role = role.replace("-", "_")

    try:
        user = get_user_with_role_check(user_id, ["team_member", "team_lead"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)

    if user.role != normalized_role:
        return HTMLResponse("Access denied", status_code=403)

    if not user.department_id:
        return HTMLResponse("No department assigned", status_code=400)

    notifications = db.query(models.Notification).filter(
        models.Notification.department_id == user.department_id,
        models.Notification.is_read == False
    ).all()

    try:
        final_assets = get_final_assets_for_department_user(user, db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    department = get_department_by_id(user.department_id, db)

    return templates.TemplateResponse(
    request,
    "final_assets_dashboard.html",
    {
        "request": request,
        "user": user,
        "department": department,
        "assets": final_assets,
        "notifications": notifications
    }
)
    
    
    
@app.get("/ui/{role}/{user_id}/final-assets/export")
def export_final_assets_for_department_user(
    role: str,
    user_id: int,
    db: Session = Depends(get_db)
):
    normalized_role = role.replace("-", "_")

    try:
        user = get_user_with_role_check(user_id, ["team_member", "team_lead"], db)
    except Exception as e:
        raise HTTPException(status_code=403, detail=str(e))

    if user.role != normalized_role:
        raise HTTPException(status_code=403, detail="Access denied")

    if not user.department_id:
        raise HTTPException(status_code=400, detail="No department assigned")

    final_assets = get_final_assets_for_department_user(user, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Final Approved Assets"

    ws.append([
        "ID",
        "Serial No",
        "Asset / Serial No",
        "Asset Type",
        "Host Name",
        "Description",
        "IP Address",
        "Privilege",
        "Criticality",
        "Condition",
        "Assigned To / Custodian / Product Owner / Business Line"
    ])

    for asset in final_assets:
        ws.append([
            asset.id,
            asset.serial_no,
            asset.asset_serial_no,
            asset.asset_type,
            asset.host_name,
            asset.description,
            asset.ip_address,
            asset.privilege,
            asset.criticality,
            asset.condition,
            asset.assigned_to_custodian_owner_business_line
        ])

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / f"final_assets_{normalized_role}_{user.department_id}.xlsx"
    wb.save(str(file_path))

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    

@app.get("/ui/final-assets/{user_id}/export-pdf")
def export_final_assets_pdf(user_id: int, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if user.role not in ["team_member", "team_lead"]:
        raise HTTPException(status_code=403, detail="Access denied")

    if not user.department_id:
        raise HTTPException(status_code=400, detail="No department assigned")

    department = db.query(models.Department).filter(
        models.Department.id == user.department_id
    ).first()

    assets = db.query(models.Asset).filter(
        models.Asset.department_id == user.department_id
    ).order_by(models.Asset.id.desc()).all()

    approved_validations = db.query(models.AssetValidation).filter(
        models.AssetValidation.department_id == user.department_id,
        models.AssetValidation.approval_stage == "approved"
    ).all()

    approved_asset_ids = {v.asset_id for v in approved_validations if v.asset_id is not None}
    final_assets = [a for a in assets if a.id in approved_asset_ids]

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / f"final_assets_department_{user.department_id}.pdf"

    c = canvas.Canvas(str(file_path), pagesize=A4)
    width, height = A4
    y = height - 50

    c.setFont("Helvetica-Bold", 16)
    c.drawString(40, y, f"Final Approved Assets - {department.name if department else 'Department'}")
    y -= 30

    c.setFont("Helvetica", 10)
    for asset in final_assets:
        line = f"ID: {asset.id} | Serial: {asset.serial_no or '-'} | Type: {asset.asset_type or '-'} | Host: {asset.host_name or '-'}"
        c.drawString(40, y, line[:110])
        y -= 16

        if y < 50:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = height - 50

    c.save()

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/pdf"
    )


@app.get("/ui/team-member/{user_id}/export")
def export_team_member(user_id: int, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["team_member"], db)
    except Exception as e:
        raise HTTPException(status_code=403, detail=str(e))

    if not user.department_id:
        raise HTTPException(status_code=400, detail="No department assigned")

    assets = db.query(models.Asset).filter(
        models.Asset.department_id == user.department_id
    ).order_by(models.Asset.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Team Member Assets"

    ws.append([
        "ID",
        "Serial No",
        "Asset / Serial No",
        "Asset Type",
        "Host Name",
        "Description",
        "IP Address",
        "Privilege",
        "Criticality",
        "Condition",
        "Assigned To / Custodian / Product Owner / Business Line"
    ])

    for asset in assets:
        ws.append([
            asset.id,
            asset.serial_no,
            asset.asset_serial_no,
            asset.asset_type,
            asset.host_name,
            asset.description,
            asset.ip_address,
            asset.privilege,
            asset.criticality,
            asset.condition,
            asset.assigned_to_custodian_owner_business_line
        ])

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / f"team_member_assets_{user.department_id}.xlsx"
    wb.save(str(file_path))

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
@app.get("/ui/team-lead/{user_id}/export")
def export_team_lead(user_id: int, db: Session = Depends(get_db)):
    try:
        user = get_user_with_role_check(user_id, ["team_lead"], db)
    except Exception as e:
        raise HTTPException(status_code=403, detail=str(e))

    if not user.department_id:
        raise HTTPException(status_code=400, detail="No department assigned")

    assets = db.query(models.Asset).filter(
        models.Asset.department_id == user.department_id
    ).order_by(models.Asset.id.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Team Lead Assets"

    ws.append([
        "ID",
        "Serial No",
        "Asset / Serial No",
        "Asset Type",
        "Host Name",
        "Description",
        "IP Address",
        "Privilege",
        "Criticality",
        "Condition",
        "Assigned To / Custodian / Product Owner / Business Line"
    ])

    for asset in assets:
        ws.append([
            asset.id,
            asset.serial_no,
            asset.asset_serial_no,
            asset.asset_type,
            asset.host_name,
            asset.description,
            asset.ip_address,
            asset.privilege,
            asset.criticality,
            asset.condition,
            asset.assigned_to_custodian_owner_business_line
        ])

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    file_path = output_dir / f"team_lead_assets_{user.department_id}.xlsx"
    wb.save(str(file_path))

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@app.get("/ui/{role}/{user_id}/export")
def export_admin_assets(role: str, user_id: int, department_id: int | None = None, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="User not found")

    if user.role not in ["compliance", "infosec"]:
        raise HTTPException(status_code=403, detail="Access denied")

    approved_validations_query = db.query(models.AssetValidation).filter(
        models.AssetValidation.approval_stage == "approved"
    )

    if department_id:
        approved_validations_query = approved_validations_query.filter(
            models.AssetValidation.department_id == department_id
        )

    approved_validations = approved_validations_query.all()
    approved_asset_ids = {v.asset_id for v in approved_validations if v.asset_id is not None}

    assets_query = db.query(models.Asset)
    if department_id:
        assets_query = assets_query.filter(models.Asset.department_id == department_id)

    assets = assets_query.order_by(models.Asset.id.desc()).all()
    final_assets = [a for a in assets if a.id in approved_asset_ids]

    department_map = {
        d.id: d.name for d in db.query(models.Department).all()
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Final Approved Assets"

    ws.append([
        "ID",
        "Serial No",
        "Asset / Serial No",
        "Asset Type",
        "Host Name",
        "Description",
        "IP Address",
        "Privilege",
        "Criticality",
        "Condition",
        "Assigned To / Custodian / Product Owner / Business Line",
        "Department"
    ])

    for asset in final_assets:
        ws.append([
            asset.id,
            asset.serial_no,
            asset.asset_serial_no,
            asset.asset_type,
            asset.host_name,
            asset.description,
            asset.ip_address,
            asset.privilege,
            asset.criticality,
            asset.condition,
            asset.assigned_to_custodian_owner_business_line,
            department_map.get(asset.department_id, "N/A")
        ])

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = f"_dept_{department_id}" if department_id else "_all"
    file_path = output_dir / f"final_assets{suffix}.xlsx"

    wb.save(str(file_path))

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
@app.get("/ui/{role}/{user_id}/final-assets-admin", response_class=HTMLResponse)
def final_assets_admin_dashboard(
    role: str,
    user_id: int,
    request: Request,
    department_id: int | None = None,
    db: Session = Depends(get_db)
):
    normalized_role = role.replace("-", "_")

    try:
        user = get_user_with_role_check(user_id, ["compliance", "infosec"], db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=403)

    if user.role != normalized_role:
        return HTMLResponse("Access denied", status_code=403)

    notifications = []
    departments = db.query(models.Department).order_by(models.Department.name).all()
    department_map = get_department_map(db)

    try:
        final_assets = get_final_assets_for_admin(normalized_role, department_id, db)
    except Exception as e:
        return HTMLResponse(str(e), status_code=400)

    return templates.TemplateResponse(
    request,
    "final_assets_admin_dashboard.html",
    {
        "request": request,
        "user": user,
        "departments": departments,
        "department_map": department_map,
        "selected_department_id": department_id,
        "assets": final_assets,
        "notifications": notifications
    }
)

@app.get("/ui/{role}/{user_id}/final-assets-admin/export")
def export_final_assets_for_admin(
    role: str,
    user_id: int,
    department_id: int | None = None,
    db: Session = Depends(get_db)
):
    normalized_role = role.replace("-", "_")

    try:
        user = get_user_with_role_check(user_id, ["compliance", "infosec"], db)
    except Exception as e:
        raise HTTPException(status_code=403, detail=str(e))

    if user.role != normalized_role:
        raise HTTPException(status_code=403, detail="Access denied")

    final_assets = get_final_assets_for_admin(normalized_role, department_id, db)
    department_map = get_department_map(db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Final Approved Assets"

    ws.append([
        "ID",
        "Serial No",
        "Asset / Serial No",
        "Asset Type",
        "Host Name",
        "Description",
        "IP Address",
        "Privilege",
        "Criticality",
        "Condition",
        "Assigned To / Custodian / Product Owner / Business Line",
        "Department"
    ])

    for asset in final_assets:
        ws.append([
            asset.id,
            asset.serial_no,
            asset.asset_serial_no,
            asset.asset_type,
            asset.host_name,
            asset.description,
            asset.ip_address,
            asset.privilege,
            asset.criticality,
            asset.condition,
            asset.assigned_to_custodian_owner_business_line,
            department_map.get(asset.department_id, "N/A")
        ])

    output_dir = Path("generated_reports")
    output_dir.mkdir(parents=True, exist_ok=True)
    suffix = f"_dept_{department_id}" if department_id else "_all"
    file_path = output_dir / f"final_assets_admin_{normalized_role}{suffix}.xlsx"
    wb.save(str(file_path))

    return FileResponse(
        path=str(file_path),
        filename=file_path.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
@app.post("/upload/assets")
def upload_assets(file: UploadFile = File(...), db: Session = Depends(get_db)):
    contents = file.file.read()

    temp_path = Path("temp_upload.xlsx")
    with open(temp_path, "wb") as f:
        f.write(contents)

    wb = load_workbook(temp_path)

    results = {
        "departments_created": 0,
        "assets_created": 0,
        "rows_skipped": 0,
        "errors": []
    }

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]

        if sheet.max_row < 2:
            continue

        department_name = sheet_name.strip()

        dept = db.query(models.Department).filter(
            models.Department.name == department_name
        ).first()

        if not dept:
            dept = models.Department(name=department_name)
            db.add(dept)
            db.commit()
            db.refresh(dept)
            results["departments_created"] += 1

        headers = [cell.value for cell in sheet[1]]

        # ✅ YOUR BLOCK GOES HERE
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row_data = dict(zip(headers, row))

            # ✅ S/N identifier
            raw_sn = row_data.get("S/N")

            if raw_sn is None or str(raw_sn).strip() == "":
                results["rows_skipped"] += 1
                results["errors"].append({
                    "sheet": sheet_name,
                    "row": row_index,
                    "field": "S/N",
                    "value": None,
                    "error": "Missing S/N identifier"
                })
                continue

            sn_value = str(raw_sn).strip()
            serial = f"{department_name}-{sn_value}"

            condition = row_data.get("Condition")
            privilege = row_data.get("Priviledge")
            criticality = row_data.get("Criticality")

            condition = str(condition).strip() if condition else None
            privilege = str(privilege).strip() if privilege else None
            criticality = str(criticality).strip() if criticality else None

            row_has_error = False

            # Validate only if present
            if condition and condition not in ALLOWED_CONDITIONS:
                results["errors"].append({
                    "sheet": sheet_name,
                    "row": row_index,
                    "field": "Condition",
                    "value": condition,
                    "error": f"Invalid value. Allowed: {sorted(ALLOWED_CONDITIONS)}"
                })
                row_has_error = True

            if privilege and privilege not in ALLOWED_PRIVILEGES:
                results["errors"].append({
                    "sheet": sheet_name,
                    "row": row_index,
                    "field": "Priviledge",
                    "value": privilege,
                    "error": f"Invalid value. Allowed: {sorted(ALLOWED_PRIVILEGES)}"
                })
                row_has_error = True

            if criticality and criticality not in ALLOWED_CRITICALITIES:
                results["errors"].append({
                    "sheet": sheet_name,
                    "row": row_index,
                    "field": "Criticality",
                    "value": criticality,
                    "error": f"Invalid value. Allowed: {sorted(ALLOWED_CRITICALITIES)}"
                })
                row_has_error = True

            if row_has_error:
                results["rows_skipped"] += 1
                continue

            existing = db.query(models.Asset).filter(
                models.Asset.serial_no == serial
            ).first()

            if existing:
                results["rows_skipped"] += 1
                results["errors"].append({
                    "sheet": sheet_name,
                    "row": row_index,
                    "field": "S/N",
                    "value": sn_value,
                    "error": f"Duplicate identifier: {serial}"
                })
                continue

            asset = models.Asset(
                serial_no=serial,
                asset_serial_no=row_data.get("Asset / Serial No") or None,
                asset_type=row_data.get("Asset Type") or None,
                host_name=row_data.get("Host / Application / Server name") or None,
                description=row_data.get("Description") or None,
                ip_address=row_data.get("Application URL / IP") or None,
                privilege=privilege,
                criticality=criticality,
                condition=condition,
                assigned_to_custodian_owner_business_line=row_data.get(
                    "Assigned to / Custodian/ Product Owner/ Business Line"
                ) or None,
                department_id=dept.id
            )

            db.add(asset)
            results["assets_created"] += 1

    db.commit()

    return {
        "message": "Upload completed",
        "summary": {
            "departments_created": results["departments_created"],
            "assets_created": results["assets_created"],
            "rows_skipped": results["rows_skipped"],
            "error_count": len(results["errors"])
        },
        "errors": results["errors"]
    }
    
    
@app.get("/ui/upload-assets/{user_id}", response_class=HTMLResponse)
def upload_assets_page(user_id: int, request: Request, db: Session = Depends(get_db)):
    user = db.query(models.ADUser).filter(models.ADUser.id == user_id).first()
    if not user:
        return HTMLResponse("User not found", status_code=404)

    if user.role not in ["compliance", "infosec"]:
        return HTMLResponse("Access denied", status_code=403)

    notifications = []

    return templates.TemplateResponse(
        "upload_assets.html",
        {
            "request": request,
            "user": user,
            "notifications": notifications
        }
    )
